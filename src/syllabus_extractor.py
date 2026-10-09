"""
Turns uploaded syllabus files into dated DigestItems using Gemini.

Syllabi live in syllabi/<COURSE_KEY>/*.{pdf,docx}, one folder per course.
COURSE_KEY is matched against the live Canvas course list (case-insensitive,
substring either direction) so it survives semester rollover automatically —
nothing here hardcodes a course id or term.

Each file is sent to Gemini once. The result is saved next to it as
<file>.extracted.json, tagged with the source file's SHA-256, and every
later run reads that JSON instead of calling the API again. Replacing the
file changes its hash and triggers a fresh extraction. The JSON is also the
correction layer: if the model misreads a date, edit the JSON on GitHub and
the fix sticks until the syllabus file itself changes.

Extraction is still a model reading prose, not a guarantee — every item is
reconciled against Canvas in merge.py, and a disagreement between the two
gets flagged rather than silently resolved.
"""

from __future__ import annotations

import hashlib
import json
import logging
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Callable, Literal

from pydantic import BaseModel, Field, ValidationError

from .digest import LOCAL_TZ, DigestItem

logger = logging.getLogger(__name__)

# Free-tier Gemini model; override with the GEMINI_MODEL repo variable.
DEFAULT_MODEL = "gemini-3.8-flash"
# Also free-tier; tried in order only when the main model is overloaded or
# not offered to this key (Google retires older models for new accounts).
FALLBACK_MODELS = ("gemini-3.7-flash", "gemini-3.5-flash")
# Per request. A 13-page syllabus normally takes well under a minute.
_REQUEST_TIMEOUT_MS = 90_000
EXTRACTED_SUFFIX = ".extracted.json"
_SOURCE_SUFFIXES = (".pdf", ".docx")

# Items with no stated time are treated as due at end of day, but flagged
# so the dashboard shows just the date instead of an invented time.
_DEFAULT_DUE_TIME = time(23, 59)


class ExtractedItem(BaseModel):
    title: str = Field(description="The syllabus's own short name for the deliverable, without the date.")
    due_date: str = Field(description="Due date as YYYY-MM-DD.")
    due_time: str | None = Field(
        default=None, description="Due time as 24-hour HH:MM, only if the syllabus states one; otherwise null."
    )
    kind: Literal["assignment", "quiz", "discussion"] = Field(
        description='"quiz" for quizzes, tests and exams; "discussion" for graded discussion posts; otherwise "assignment".'
    )


class ExtractionResult(BaseModel):
    items: list[ExtractedItem]


# (path, course_name, term_description, today) -> (extracted items, model that produced them)
ExtractFn = Callable[[Path, str, str, date], tuple[list[ExtractedItem], str]]


@dataclass
class FileStatus:
    folder: str
    filename: str
    state: str  # "cached" | "extracted" | "pending" | "failed" | "unmatched"
    item_count: int = 0
    detail: str = ""


_PROMPT = """\
You are extracting graded deliverables from a college course syllabus for a \
student's assignment tracker.

Course: {course_name}
Term: {term}
Today's date: {today}

Return every graded item that has a specific calendar due date: assignments, \
homework, papers, projects and project milestones, labs, quizzes, exams, \
presentations, and graded discussion posts.

Do not include class topics, reading schedules, holidays, "no class" days, \
office hours, add/drop or withdrawal deadlines, or anything described only by \
a recurring pattern (such as "every Friday") without specific dates.

- If a date has no year, use the year that places it within the course term.
- If one schedule row lists several deliverables, return each one separately.
- Keep titles short and use the syllabus's own wording.
If the syllabus lists no dated deliverables, return an empty list.
"""


def _normalize(s: str) -> str:
    return "".join(ch for ch in s.lower() if ch.isalnum())


def match_course_key(folder_name: str, courses: list[dict]) -> int | None:
    """Matches a syllabi/<folder_name> against the live Canvas course list."""
    key = _normalize(folder_name)
    if not key:
        return None
    for course in courses:
        code = _normalize(course.get("course_code") or "")
        if code and (key in code or code in key):
            return course.get("id")
    for course in courses:
        name = _normalize(course.get("name") or "")
        if name and (key in name or name in key):
            return course.get("id")
    return None


def discover_syllabus_files(root: Path) -> dict[str, list[Path]]:
    """Maps each immediate subfolder of `root` to its .pdf/.docx files."""
    if not root.is_dir():
        return {}
    manifest: dict[str, list[Path]] = {}
    for folder in sorted(p for p in root.iterdir() if p.is_dir()):
        files = sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in _SOURCE_SUFFIXES)
        if files:
            manifest[folder.name] = files
    return manifest


def sidecar_path(source: Path) -> Path:
    return source.with_name(source.name + EXTRACTED_SUFFIX)


def remove_orphaned_extractions(root: Path) -> list[Path]:
    """Deletes .extracted.json files whose syllabus file has been removed."""
    if not root.is_dir():
        return []
    removed: list[Path] = []
    for sidecar in root.glob(f"*/*{EXTRACTED_SUFFIX}"):
        source = sidecar.with_name(sidecar.name[: -len(EXTRACTED_SUFFIX)])
        if not source.exists():
            sidecar.unlink()
            removed.append(sidecar)
    return removed


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _docx_text(path: Path) -> str:
    """Paragraphs plus tables — syllabus schedules are usually tables, which
    python-docx's `paragraphs` alone skips entirely."""
    from docx import Document

    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def describe_term(course: dict) -> str:
    term = course.get("term") or {}
    name = term.get("name") or "unknown"
    start, end = (term.get("start_at") or "")[:10], (term.get("end_at") or "")[:10]
    return f"{name} ({start} to {end})" if start and end else name


def make_gemini_extractor(api_key: str, model: str = DEFAULT_MODEL) -> ExtractFn:
    import httpx
    from google import genai
    from google.genai import errors, types

    # Free-tier requests are the first to get 503 "high demand" or 429
    # rate-limit responses, so back off and retry well past the SDK default.
    # The SDK has no timeout by default, so an overloaded server that never
    # answers would otherwise hang the run indefinitely.
    client = genai.Client(
        api_key=api_key,
        http_options=types.HttpOptions(
            timeout=_REQUEST_TIMEOUT_MS,
            retry_options=types.HttpRetryOptions(
                attempts=3, initial_delay=5, max_delay=60, http_status_codes=[429, 500, 503, 504]
            ),
        ),
    )
    # If the chosen model stays overloaded, another free model usually isn't.
    models = [model] + [m for m in FALLBACK_MODELS if m != model]

    def extract(path: Path, course_name: str, term: str, today: date) -> tuple[list[ExtractedItem], str]:
        if path.suffix.lower() == ".pdf":
            # Sent as a PDF, not extracted text: Gemini reads the page
            # layout, which keeps schedule tables intact.
            document = types.Part.from_bytes(data=path.read_bytes(), mime_type="application/pdf")
        else:
            document = types.Part.from_text(text=_docx_text(path))
        prompt = _PROMPT.format(course_name=course_name, term=term, today=today.isoformat())
        config = types.GenerateContentConfig(
            response_mime_type="application/json",
            response_json_schema=ExtractionResult.model_json_schema(),
            automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),
        )

        for i, current in enumerate(models):
            logger.info("Sending %s to %s...", path.name, current)
            try:
                response = client.models.generate_content(model=current, contents=[document, prompt], config=config)
            except (errors.APIError, httpx.TimeoutException) as e:
                if isinstance(e, errors.APIError):
                    reason = str(e.code)
                    # 404: the model was retired or isn't offered to this key.
                    overloaded = e.code in (404, 429) or (e.code or 0) >= 500
                else:
                    reason, overloaded = "timed out", True
                if overloaded and i + 1 < len(models):
                    logger.warning("%s unavailable (%s); trying %s", current, reason, models[i + 1])
                    continue
                raise
            if not response.text:
                raise ValueError(f"{current} returned an empty response")
            return ExtractionResult.model_validate_json(response.text).items, current
        raise AssertionError("unreachable")

    return extract


def _load_sidecar(sidecar: Path, source_hash: str) -> list[ExtractedItem] | None:
    """Saved items if the sidecar matches the current file, else None.

    Raises ValueError for a sidecar that exists but can't be read, so a
    broken hand edit is reported instead of silently overwritten.
    """
    if not sidecar.exists():
        return None
    try:
        data = json.loads(sidecar.read_text(encoding="utf-8"))
        if data.get("source_sha256") != source_hash:
            return None
        return ExtractionResult.model_validate({"items": data.get("items", [])}).items
    except (json.JSONDecodeError, ValidationError, AttributeError) as e:
        raise ValueError(f"{sidecar.name} is not valid ({e.__class__.__name__}); fix it or delete it to re-extract") from e


def _write_sidecar(sidecar: Path, source_hash: str, model: str, items: list[ExtractedItem]) -> None:
    data = {
        "source_sha256": source_hash,
        "model": model,
        "extracted_at": datetime.now(tz=LOCAL_TZ).isoformat(timespec="seconds"),
        "items": [item.model_dump() for item in items],
    }
    sidecar.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _to_digest_item(item: ExtractedItem, course_name: str) -> DigestItem | None:
    try:
        due_date = date.fromisoformat(item.due_date)
        due_time = time.fromisoformat(item.due_time) if item.due_time else None
    except ValueError:
        logger.warning("Skipping syllabus item %r with unreadable date %r / time %r", item.title, item.due_date, item.due_time)
        return None
    return DigestItem(
        course_name=course_name,
        title=item.title,
        due_at=datetime.combine(due_date, due_time or _DEFAULT_DUE_TIME, tzinfo=LOCAL_TZ),
        html_url="",
        points_possible=None,
        kind=item.kind,
        source="Syllabus",
        has_time=due_time is not None,
    )


def process_syllabi(
    manifest: dict[str, list[Path]],
    courses: list[dict],
    *,
    extract_fn: ExtractFn | None,
    today: date | None = None,
) -> tuple[list[DigestItem], list[FileStatus]]:
    """
    Loads saved extractions, extracting (and saving) any file that has none.

    `extract_fn` is None when no API key is configured: saved extractions
    are still used, and new files are reported as pending.
    """
    today = today or datetime.now(tz=LOCAL_TZ).date()
    results: list[DigestItem] = []
    statuses: list[FileStatus] = []

    for folder_name, paths in manifest.items():
        course_id = match_course_key(folder_name, courses)
        if course_id is None:
            logger.warning("No Canvas course matches syllabus folder %r - skipping.", folder_name)
            statuses.extend(FileStatus(folder_name, p.name, "unmatched") for p in paths)
            continue
        course = next(c for c in courses if c.get("id") == course_id)
        course_name = course.get("name") or course.get("course_code") or f"Course {course_id}"

        for path in paths:
            status = FileStatus(folder_name, path.name, "cached")
            sidecar = sidecar_path(path)
            source_hash = _sha256(path)
            try:
                items = _load_sidecar(sidecar, source_hash)
                if items is None:
                    if extract_fn is None:
                        status.state = "pending"
                        statuses.append(status)
                        continue
                    items, used_model = extract_fn(path, course_name, describe_term(course), today)
                    _write_sidecar(sidecar, source_hash, used_model, items)
                    status.state = "extracted"
            except Exception as e:
                # One bad file (rate limit, malformed PDF, broken hand edit)
                # shouldn't drop every other syllabus. Nothing is saved, so
                # an API failure is retried on the next run.
                logger.warning("Syllabus %s/%s failed: %s", folder_name, path.name, e)
                statuses.append(FileStatus(folder_name, path.name, "failed", detail=str(e)[:200]))
                continue

            converted = [d for d in (_to_digest_item(i, course_name) for i in items) if d is not None]
            status.item_count = len(converted)
            statuses.append(status)
            results.extend(converted)

    return results, statuses


if __name__ == "__main__":
    # Try extraction on one file without Canvas or the rest of the pipeline:
    #   python -m src.syllabus_extractor syllabi/COSC240/syllabus.pdf
    # Or list the models this key can use:
    #   python -m src.syllabus_extractor --list-models
    # Reads GEMINI_API_KEY (and optionally GEMINI_MODEL) from the environment.
    # Prints what was found; writes nothing.
    import os
    import sys

    from .envfile import load_env_file

    load_env_file()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    for noisy in ("httpx", "google_genai"):
        logging.getLogger(noisy).setLevel(logging.WARNING)
    if len(sys.argv) != 2:
        sys.exit("usage: python -m src.syllabus_extractor <path-to-syllabus.pdf|.docx> | --list-models")
    api_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if not api_key:
        sys.exit("GEMINI_API_KEY is not set (in this terminal or in .env).")

    if sys.argv[1] == "--list-models":
        from google import genai

        client = genai.Client(api_key=api_key)  # must outlive the paged listing
        for m in client.models.list():
            actions = ", ".join(m.supported_actions or [])
            print(f"{m.name.removeprefix('models/'):45} {actions}")
        sys.exit(0)
    source = Path(sys.argv[1])
    if not source.is_file():
        sys.exit(f"No such file: {source}")

    model = os.environ.get("GEMINI_MODEL", "").strip() or DEFAULT_MODEL
    today = datetime.now(tz=LOCAL_TZ).date()
    found, used_model = make_gemini_extractor(api_key, model)(source, source.parent.name, "unknown", today)
    print(f"{len(found)} item(s) from {source.name} using {used_model}:")
    for item in found:
        print(f"  {item.due_date}  {item.due_time or '--:--'}  [{item.kind}]  {item.title}")
